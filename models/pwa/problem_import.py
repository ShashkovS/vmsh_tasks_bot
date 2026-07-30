"""Pure parser and dry-run comparison for the legacy problem workbook."""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from helpers.consts import ANS_TYPES_DECODER, ANS_TYPE, PROB_TYPES_DECODER, PROB_TYPE


SHEETS = ("Задачи", "Старые")
COLUMNS = (
    "level",
    "lesson",
    "prob",
    "item",
    "title",
    "prob_text",
    "prob_type",
    "ans_type",
    "ans_validation",
    "validation_error",
    "cor_ans",
    "cor_ans_checker",
    "wrong_ans",
    "congrat",
)
MAX_ROWS = 5_000
MAX_CELL_LENGTH = 100_000


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip().replace('""', '"')
    return result or None


def _integer(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = _text(value)
    if text is None or re.fullmatch(r"[-+]?\d+", text) is None:
        return None
    return int(text)


def _decoded(value: object, mapping: dict[str, Any]) -> int | None:
    text = _text(value)
    if text is None:
        return None
    folded = text.casefold()
    for label, decoded in mapping.items():
        if label.casefold() == folded:
            return int(decoded)
    return None


def _diagnostic(sheet: str, row: int, field: str, code: str) -> dict[str, object]:
    return {"sheet": sheet, "row": row, "field": field, "code": code}


def parse_problem_workbook(
    source: bytes,
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    """Parse the two legacy sheets without executing formulas or touching SQLite."""

    try:
        workbook = load_workbook(
            io.BytesIO(source), read_only=True, data_only=False, keep_links=False
        )
    except (
        InvalidFileException,
        OSError,
        ValueError,
        EOFError,
        zipfile.BadZipFile,
    ) as error:
        raise ValueError("invalid_workbook") from error

    present = [name for name in SHEETS if name in workbook.sheetnames]
    if not present:
        raise ValueError("missing_problem_sheets")

    rows: list[dict[str, object]] = []
    diagnostics: list[dict[str, object]] = []
    for sheet_name in present:
        sheet = workbook[sheet_name]
        header = tuple(_text(cell) for cell in next(sheet.iter_rows(values_only=True)))
        if header[: len(COLUMNS)] != COLUMNS:
            diagnostics.append(_diagnostic(sheet_name, 1, "header", "invalid_header"))
            continue
        # Row 2 contains human-readable column help in the canonical workbook.
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=3, values_only=True), 3
        ):
            values = cells[: len(COLUMNS)]
            if all(_text(value) is None for value in values[:4]):
                continue
            if len(rows) >= MAX_ROWS:
                raise ValueError("too_many_rows")
            row_diagnostics: list[dict[str, object]] = []
            for index, value in enumerate(values):
                text = _text(value)
                if text is not None and len(text) > MAX_CELL_LENGTH:
                    row_diagnostics.append(
                        _diagnostic(
                            sheet_name, row_number, COLUMNS[index], "cell_too_long"
                        )
                    )

            group_code = _text(values[0])
            lesson = _integer(values[1])
            problem = _integer(values[2])
            item = _text(values[3]) or ""
            title = _text(values[4])
            problem_type = _decoded(values[6], PROB_TYPES_DECODER)
            answer_type = _decoded(values[7], ANS_TYPES_DECODER)
            answer_validation = _text(values[8])

            if group_code is None:
                row_diagnostics.append(
                    _diagnostic(sheet_name, row_number, "level", "group_required")
                )
            if lesson is None or lesson < 0:
                row_diagnostics.append(
                    _diagnostic(sheet_name, row_number, "lesson", "lesson_invalid")
                )
            if problem is None or problem <= 0:
                row_diagnostics.append(
                    _diagnostic(
                        sheet_name, row_number, "prob", "problem_number_invalid"
                    )
                )
            if title is None:
                row_diagnostics.append(
                    _diagnostic(sheet_name, row_number, "title", "title_required")
                )
            if problem_type is None:
                row_diagnostics.append(
                    _diagnostic(
                        sheet_name, row_number, "prob_type", "problem_type_invalid"
                    )
                )
            if problem_type == int(PROB_TYPE.TEST) and answer_type is None:
                row_diagnostics.append(
                    _diagnostic(
                        sheet_name, row_number, "ans_type", "answer_type_invalid"
                    )
                )
            if problem_type != int(PROB_TYPE.TEST):
                answer_type = None
            if answer_validation is not None and answer_type not in {
                None,
                int(ANS_TYPE.SELECT_ONE),
            }:
                try:
                    re.compile(answer_validation)
                except re.error:
                    row_diagnostics.append(
                        _diagnostic(
                            sheet_name,
                            row_number,
                            "ans_validation",
                            "answer_validation_invalid",
                        )
                    )

            rows.append(
                {
                    "sheet": sheet_name,
                    "row": row_number,
                    "group_code": group_code,
                    "lesson": lesson,
                    "problem": problem,
                    "item": item,
                    "title": title,
                    "problem_text": _text(values[5]) or "",
                    "problem_type": problem_type,
                    "answer_type": answer_type,
                    "answer_validation": answer_validation,
                    "validation_error": _text(values[9]),
                    "correct_answer": _text(values[10]),
                    "correct_answer_checker": _text(values[11]),
                    "wrong_answer": _text(values[12]),
                    "congratulation": _text(values[13]),
                    "diagnostics": row_diagnostics,
                }
            )
            diagnostics.extend(row_diagnostics)

    keys = [
        (str(row["group_code"]).casefold(), row["lesson"], row["problem"], row["item"])
        for row in rows
        if row["group_code"] is not None
        and row["lesson"] is not None
        and row["problem"] is not None
    ]
    duplicate_keys = {key for key, count in Counter(keys).items() if count > 1}
    for row in rows:
        key = (
            str(row["group_code"]).casefold(),
            row["lesson"],
            row["problem"],
            row["item"],
        )
        if key in duplicate_keys:
            item = _diagnostic(
                str(row["sheet"]), int(row["row"]), "prob", "duplicate_problem"
            )
            row["diagnostics"].append(item)  # type: ignore[union-attr]
            diagnostics.append(item)
    return rows, diagnostics


def compare_problem_rows(
    rows: list[dict[str, object]],
    groups: list[dict[str, object]],
    current: list[dict[str, object]],
) -> list[dict[str, object]]:
    """Add group identity and create/update/unchanged action to parsed rows."""

    groups_by_code = {str(group["short_code"]).casefold(): group for group in groups}
    current_by_key = {
        (
            str(problem["group_id"]),
            int(problem["lesson"]),
            int(problem["prob"]),
            str(problem["item"]),
        ): problem
        for problem in current
    }
    compared: list[dict[str, object]] = []
    fields = (
        "title",
        "problem_text",
        "problem_type",
        "answer_type",
        "answer_validation",
        "validation_error",
        "correct_answer",
        "correct_answer_checker",
        "wrong_answer",
        "congratulation",
    )
    database_fields = {
        "problem_text": "prob_text",
        "problem_type": "prob_type",
        "answer_type": "ans_type",
        "answer_validation": "ans_validation",
        "correct_answer": "cor_ans",
        "correct_answer_checker": "cor_ans_checker",
        "wrong_answer": "wrong_ans",
        "congratulation": "congrat",
    }
    numeric_fields = {"problem_type", "answer_type"}
    for parsed in rows:
        row = {
            **parsed,
            "group_id": None,
            "group_public_id": None,
            "problem_public_id": None,
        }
        group_code = parsed["group_code"]
        group = groups_by_code.get(str(group_code).casefold()) if group_code else None
        if group is None:
            diagnostic = _diagnostic(
                str(parsed["sheet"]), int(parsed["row"]), "level", "group_unknown"
            )
            row["diagnostics"] = [*parsed["diagnostics"], diagnostic]  # type: ignore[misc]
        else:
            row["group_id"] = group["group_id"]
            row["group_public_id"] = group["public_id"]
        if row["diagnostics"]:
            row["action"] = "invalid"
            compared.append(row)
            continue
        key = (
            str(row["group_id"]),
            int(row["lesson"]),
            int(row["problem"]),
            str(row["item"]),
        )
        stored = current_by_key.get(key)
        if stored is None:
            row["action"] = "create"
        else:
            row["problem_public_id"] = stored.get("public_id")
            same = True
            for field in fields:
                stored_value = stored.get(database_fields.get(field, field))
                if field not in numeric_fields:
                    stored_value = _text(stored_value)
                    if field == "problem_text":
                        stored_value = stored_value or ""
                if row[field] != stored_value:
                    same = False
                    break
            row["action"] = "unchanged" if same else "update"
        compared.append(row)
    return compared


def problem_import_preview_hash(
    course_public_id: str,
    source_sha256: str,
    rows: list[dict[str, object]],
) -> str:
    payload = [
        {
            key: row[key]
            for key in (
                "sheet",
                "row",
                "group_code",
                "group_public_id",
                "lesson",
                "problem",
                "item",
                "title",
                "problem_text",
                "problem_type",
                "answer_type",
                "answer_validation",
                "validation_error",
                "correct_answer",
                "correct_answer_checker",
                "wrong_answer",
                "congratulation",
                "action",
                "problem_public_id",
                "diagnostics",
            )
        }
        for row in rows
    ]
    encoded = json.dumps(
        {"course": course_public_id, "source": source_sha256, "rows": payload},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode()
    return hashlib.sha256(encoded).hexdigest()


__all__ = [
    "compare_problem_rows",
    "parse_problem_workbook",
    "problem_import_preview_hash",
]
