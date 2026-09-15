"""Read the historical `IDd / Уровень / Аудитория` Excel export."""

from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import load_workbook

from models.pwa.classrooms import prepare_classroom_name


REQUIRED_HEADERS = ("IDd", "Уровень", "Аудитория")


class ClassroomImportSourceError(ValueError):
    pass


def _student_id(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int) and value > 0:
        return value
    if isinstance(value, float) and value.is_integer() and value > 0:
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        parsed = int(value.strip())
        return parsed if parsed > 0 else None
    return None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_classroom_export(
    path: str | Path, *, sheet_name: str = "Итог"
) -> tuple[str, int, list[dict[str, object]]]:
    """Return source hash, header row and normalized source rows."""

    source = Path(path)
    before = source.read_bytes()
    source_sha256 = hashlib.sha256(before).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ClassroomImportSourceError(f"Worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        header_row = 0
        columns: dict[str, int] = {}
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1
        ):
            candidates = {_text(value): index for index, value in enumerate(values)}
            if all(header in candidates for header in REQUIRED_HEADERS):
                header_row = row_number
                columns = {header: candidates[header] for header in REQUIRED_HEADERS}
                break
        if not header_row:
            raise ClassroomImportSourceError(
                "The worksheet has no IDd / Уровень / Аудитория header row"
            )

        rows: list[dict[str, object]] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            raw_id = values[columns["IDd"]] if columns["IDd"] < len(values) else None
            raw_group = (
                values[columns["Уровень"]] if columns["Уровень"] < len(values) else None
            )
            raw_room = (
                values[columns["Аудитория"]]
                if columns["Аудитория"] < len(values)
                else None
            )
            if raw_id is None and raw_group is None and raw_room is None:
                continue
            room_text = _text(raw_room)
            room_name = ""
            room_key = ""
            if room_text:
                room_name, room_key = prepare_classroom_name(room_text)
            rows.append(
                {
                    "row_number": row_number,
                    "user_id": _student_id(raw_id),
                    "raw_user_id": _text(raw_id),
                    "group_label": _text(raw_group),
                    "room_name": room_name,
                    "room_key": room_key,
                }
            )
    finally:
        workbook.close()

    if hashlib.sha256(source.read_bytes()).hexdigest() != source_sha256:
        raise ClassroomImportSourceError("The Excel file changed while it was read")
    return source_sha256, header_row, rows


__all__ = [
    "ClassroomImportSourceError",
    "REQUIRED_HEADERS",
    "read_classroom_export",
]
