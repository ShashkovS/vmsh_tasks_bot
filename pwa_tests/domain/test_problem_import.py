from __future__ import annotations

from datetime import datetime, time
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from models.pwa.problem_import import (
    COLUMNS,
    compare_problem_rows,
    find_problem_import_synonym_candidates,
    parse_problem_workbook,
    problem_import_preview_hash,
)


def _workbook(*, current_rows=(), old_rows=()) -> bytes:
    workbook = Workbook()
    current = workbook.active
    current.title = "Задачи"
    old = workbook.create_sheet("Старые")
    for sheet, rows in ((current, current_rows), (old, old_rows)):
        sheet.append(COLUMNS)
        sheet.append(["Описание"] * len(COLUMNS))
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _row(
    *,
    group="н",
    lesson=41,
    problem=1,
    item="",
    title="Орехи",
    kind="Тест",
    answer="Натуральное",
):
    return (
        group,
        lesson,
        problem,
        item,
        title,
        None,
        kind,
        answer,
        None,
        "Введите число",
        29,
        None,
        "Нет",
        "Да",
    )


def test_parser_preserves_legacy_fields_and_ignores_answer_type_for_written():
    rows, diagnostics = parse_problem_workbook(
        _workbook(
            current_rows=(
                _row(),
                _row(
                    problem=2, title="Доказательство", kind="Письменно", answer="Целое"
                ),
            )
        )
    )

    assert diagnostics == []
    assert rows[0]["correct_answer"] == "29"
    assert rows[0]["problem_type"] == 1
    assert rows[0]["answer_type"] == 2
    assert rows[1]["problem_type"] == 2
    assert rows[1]["answer_type"] is None


def test_parser_preserves_excel_display_text_for_answers():
    workbook = Workbook()
    current = workbook.active
    current.title = "Задачи"
    old = workbook.create_sheet("Старые")
    for sheet in (current, old):
        sheet.append(COLUMNS)
        sheet.append(["Описание"] * len(COLUMNS))

    current.append(_row(problem=1, answer="Действительное"))
    current["K3"] = datetime(2026, 5, 2)
    current["K3"].number_format = r"d\.m"
    current.append(_row(problem=2, answer="МножЦелых"))
    current["K4"] = 7.9
    current.append(_row(problem=3, answer="Время"))
    current["K5"] = time(7, 15)
    current["K5"].number_format = "h:mm"
    current.append(_row(problem=4, answer="Дата"))
    current["K6"] = datetime(2026, 5, 29)
    current["K6"].number_format = r"dd\.mm\.yyyy"
    output = BytesIO()
    workbook.save(output)

    rows, diagnostics = parse_problem_workbook(output.getvalue())

    assert diagnostics == []
    assert [row["correct_answer"] for row in rows] == [
        "2.5",
        "7,9",
        "7:15",
        "29.05.2026",
    ]


def test_parser_reports_duplicate_identity_and_invalid_custom_regex():
    invalid = list(_row())
    invalid[8] = "["
    rows, diagnostics = parse_problem_workbook(
        _workbook(current_rows=(invalid,), old_rows=(_row(),))
    )

    assert len(rows) == 2
    assert {item["code"] for item in diagnostics} == {
        "answer_validation_invalid",
        "duplicate_problem",
    }
    assert all(row["diagnostics"] for row in rows)


def test_comparison_marks_create_update_unchanged_and_unknown_group():
    rows, _ = parse_problem_workbook(
        _workbook(
            current_rows=(
                _row(problem=1),
                _row(problem=2, title="Новое название"),
                _row(group="чужая", problem=3),
            )
        )
    )
    groups = [
        {"group_id": "beginner", "public_id": "group-beginner", "short_code": "н"}
    ]
    current = [
        {
            "group_id": "beginner",
            "lesson": 41,
            "prob": 1,
            "item": "",
            "public_id": "problem-one",
            "title": "Орехи",
            "prob_text": "",
            "prob_type": 1,
            "ans_type": 2,
            "ans_validation": "",
            "validation_error": "Введите число",
            "cor_ans": "29",
            "cor_ans_checker": "",
            "wrong_ans": "Нет",
            "congrat": "Да",
        },
        {
            "group_id": "beginner",
            "lesson": 41,
            "prob": 2,
            "item": "",
            "public_id": "problem-two",
            "title": "Старое название",
            "prob_text": "",
            "prob_type": 1,
            "ans_type": 2,
            "ans_validation": None,
            "validation_error": "Введите число",
            "cor_ans": "29",
            "cor_ans_checker": None,
            "wrong_ans": "Нет",
            "congrat": "Да",
        },
    ]

    compared = compare_problem_rows(rows, groups, current)
    assert [row["action"] for row in compared] == ["unchanged", "update", "invalid"]
    assert compared[0]["problem_public_id"] == "problem-one"
    assert compared[2]["diagnostics"][-1]["code"] == "group_unknown"

    first = problem_import_preview_hash("course-math", "a" * 64, compared)
    assert first == problem_import_preview_hash("course-math", "a" * 64, compared)
    assert first != problem_import_preview_hash("course-math", "b" * 64, compared)


def test_comparison_ignores_legacy_blank_answer_type_and_line_endings():
    written = list(
        _row(
            problem=2,
            title="Доказательство",
            kind="Письменно",
            answer=None,
        )
    )
    written[8:] = [None] * 6
    rows, _ = parse_problem_workbook(_workbook(current_rows=(written,)))
    rows[0]["validation_error"] = "Первая строка\nВторая строка"
    groups = [
        {"group_id": "beginner", "public_id": "group-beginner", "short_code": "н"}
    ]
    current = [
        {
            "group_id": "beginner",
            "lesson": 41,
            "prob": 2,
            "item": "",
            "public_id": "problem-written",
            "title": "Доказательство",
            "prob_text": "",
            "prob_type": 2,
            "ans_type": "",
            "ans_validation": None,
            "validation_error": "Первая строка\r\nВторая строка",
            "cor_ans": None,
            "cor_ans_checker": None,
            "wrong_ans": None,
            "congrat": None,
        }
    ]

    assert compare_problem_rows(rows, groups, current)[0]["action"] == "unchanged"


def test_synonym_candidates_require_equal_titles_in_different_groups():
    rows, _ = parse_problem_workbook(
        _workbook(
            current_rows=(
                _row(group="н", problem=1, title="  Орехи   и клетки "),
                _row(group="п", problem=4, title="ОРЕХИ И КЛЕТКИ"),
                _row(group="н", problem=2, title="Другая задача"),
            )
        )
    )
    groups = [
        {"group_id": "beginner", "public_id": "group-beginner", "short_code": "н"},
        {
            "group_id": "continuing",
            "public_id": "group-continuing",
            "short_code": "п",
        },
    ]

    candidates = find_problem_import_synonym_candidates(
        compare_problem_rows(rows, groups, [])
    )

    assert len(candidates) == 1
    assert candidates[0]["lesson"] == 41
    assert candidates[0]["normalized_title"] == "орехи и клетки"
    assert candidates[0]["has_group_conflict"] is False
    assert [member["problem"] for member in candidates[0]["members"]] == [1, 4]


def test_synonym_candidate_marks_ambiguous_duplicate_inside_one_group():
    rows, _ = parse_problem_workbook(
        _workbook(
            current_rows=(
                _row(group="н", problem=1, title="Метрик"),
                _row(group="н", problem=2, title="Метрик"),
                _row(group="п", problem=3, title="Метрик"),
            )
        )
    )
    groups = [
        {"group_id": "beginner", "public_id": "group-beginner", "short_code": "н"},
        {
            "group_id": "continuing",
            "public_id": "group-continuing",
            "short_code": "п",
        },
    ]

    candidates = find_problem_import_synonym_candidates(
        compare_problem_rows(rows, groups, [])
    )

    assert candidates[0]["has_group_conflict"] is True


def test_reference_workbook_remains_a_clean_characterization_input():
    source = Path(
        "_external_pipelines/ВМШ 2025-26, информация для бота ВМШ — prod.xlsx"
    )
    rows, diagnostics = parse_problem_workbook(source.read_bytes())

    assert diagnostics == []
    assert sum(row["sheet"] == "Задачи" for row in rows) == 816
    assert sum(row["sheet"] == "Старые" for row in rows) == 997
    assert (
        next(
            row["correct_answer"]
            for row in rows
            if row["sheet"] == "Задачи" and row["row"] == 23
        )
        == "2.5"
    )
